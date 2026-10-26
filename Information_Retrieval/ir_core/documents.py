"""Ingestion: turn uploaded files / pasted text into Documents made of a tree of Nodes (structure).

A Node is a structural component (section, subsection, paragraph, proviso, heading).
Its `text` is its OWN text; its region is the span of text it owns.
"""
from __future__ import annotations
import io, json, re
from dataclasses import dataclass, field
from typing import List, Optional

META_KEYS = {"id", "title", "law", "section", "subsection", "source type", "tax year", "effective date",
             "author", "year", "venue", "source", "date", "publication date", "topic", "language",
             "jurisdiction", "category"}
RE_META = re.compile(r"^([A-Za-z ]{2,20}):\s*(.+)$")
RE_MD = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")
RE_PART = re.compile(r"^(?:Part|Chapter)\s+[IVXLC\d]+\b.*$", re.I)
RE_SEC = re.compile(r"^(?:Section|Sec\.?|Article|Rule|Clause)\s+(\d+[A-Z]?)\b[:.\-—\s]*(.*)$", re.I)
RE_SUB = re.compile(r"^\((\d+[A-Z]?)\)\s+(.*)$")
RE_PAR = re.compile(r"^\(([a-z]{1,3})\)\s+(.*)$")
RE_PROV = re.compile(r"^(Provided|Explanation)\b", re.I)
RE_NUM = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+([A-Z][^.]*)$")
RE_REF = re.compile(r"\b[Ss]ections?\s+(\d+[A-Za-z]?)")


@dataclass
class Node:
    id: str
    doc_id: str
    level: int
    type: str
    heading: str = ""
    text: str = ""
    parent: Optional[str] = None
    children: List[str] = field(default_factory=list)
    citation: str = ""
    path: str = ""
    page_start: Optional[int] = None
    page_end: Optional[int] = None


@dataclass
class Document:
    id: str
    title: str
    meta: dict
    nodes: List[Node]
    hint: str = ""

    def meta_text(self):
        return " ".join(f"{k} {v}" for k, v in self.meta.items())

    def body(self):
        return "\n".join((n.heading + " " + n.text).strip() for n in self.nodes)

    def assign(self, doc_id: str):
        ren = lambda s: s.replace("~#", doc_id + "#", 1) if s else s
        self.id = doc_id
        for n in self.nodes:
            n.id, n.doc_id, n.parent = ren(n.id), doc_id, ren(n.parent)
            n.children = [ren(c) for c in n.children]


class _Builder:
    def __init__(self, title, meta):
        self.doc = Document("~", title, meta, [])
        self.root = self._new(0, "document", title, "", None)
        self.stack = [self.root]

    def _new(self, level, typ, heading, text, parent):
        n = Node(f"~#{len(self.doc.nodes)}", "~", level, typ, heading, text, parent.id if parent else None)
        self.doc.nodes.append(n)
        if parent:
            parent.children.append(n.id)
        return n

    def add(self, level, typ, heading, text="", label=None):
        while self.stack[-1].level >= level:
            self.stack.pop()
        parent = self.stack[-1]
        n = self._new(level, typ, heading, text, parent)
        base = parent.citation or self.doc.meta.get("Section", "")
        if typ == "section":
            n.citation = label or ""
        elif typ in ("subsection", "paragraph"):
            n.citation = f"{base}({label})"
            n.heading = n.citation
        elif typ == "proviso":
            n.citation = parent.citation
        n.path = (parent.path + " › " if parent.path else "") + n.heading[:70]
        self.stack.append(n)
        return n

    def text(self, s):
        cur = self.stack[-1]
        cur.text = (cur.text + " " + s).strip()


def classify(s: str):
    short = len(s) < 110 and not s.endswith(".")
    m = RE_MD.match(s)
    if m:
        return len(m.group(1)), "heading", m.group(2), "", None
    if short and RE_PART.match(s):
        return 1, "part", s, "", None
    m = RE_SEC.match(s) if short else None
    if m:
        return 2, "section", s, "", m.group(1).upper()
    m = RE_SUB.match(s)
    if m:
        return 3, "subsection", "", m.group(2), m.group(1)
    m = RE_PAR.match(s)
    if m:
        return 4, "paragraph", "", m.group(2), m.group(1)
    if RE_PROV.match(s):
        return 4, "proviso", "Proviso", s, None
    m = RE_NUM.match(s) if short else None
    if m:
        return m.group(1).count(".") + 1, "heading", s, "", None
    return None


def parse_text(text: str, name: str = "document") -> Document:
    lines = text.replace("\r", "").split("\n")
    meta, title, hint, i = {}, None, "", 0
    while i < len(lines):                       # leading "Key: value" metadata block
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        m = RE_META.match(s)
        if m and m.group(1).strip().lower() in META_KEYS:
            k, v = m.group(1).strip().lower(), m.group(2).strip()
            if k == "title":
                title = v
            elif k == "id":
                hint = v
            else:
                meta[k.title()] = v
            i += 1
            continue
        break
    body = lines[i:]
    if not title:
        for j, ln in enumerate(body):
            if re.match(r"^#\s+\S", ln.strip()):
                title = ln.strip().lstrip("#").strip()
                body = body[:j] + body[j + 1:]
                break
    b = _Builder(title or name, meta)
    b.doc.hint = hint
    for ln in body:
        s = ln.strip()
        if not s:
            continue
        c = classify(s)
        if c is None:
            b.text(s)
        else:
            level, typ, heading, rest, label = c
            b.add(level, typ, heading, rest, label)
    return b.doc


def from_json(obj, name="json") -> List[Document]:
    items = obj if isinstance(obj, list) else obj.get("documents", [obj])
    docs, reserved = [], {"title", "meta", "id", "text", "sections"}
    for d in items:
        if isinstance(d, str):
            d = {"text": d}
        b = _Builder(d.get("title") or d.get("id") or name,
                     {str(k): str(v) for k, v in (d.get("meta") or {}).items()})
        b.doc.hint = str(d.get("id", ""))
        if d.get("text"):
            b.text(d["text"])

        def walk(secs, level):
            for s in secs:
                b.add(level, "section", s.get("heading", ""), s.get("text", ""))
                walk(s.get("children", []), level + 1)
        if d.get("sections"):
            walk(d["sections"], 1)
        else:
            for k, v in d.items():
                if k not in reserved and isinstance(v, str):
                    b.add(1, "section", k.replace("_", " ").title(), v)
        docs.append(b.doc)
    return docs


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("PDF support needs: pip install pypdf")
    return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)


ITO_SECTION = re.compile(r"^\s*(?:\d+\s*\[\s*)?(\d+[A-Z]{0,3})\.\s+(.+?)\s*(?:[. ]*[—–]\s*(.*))?\s*$", re.I)
ITO_CHAPTER = re.compile(r"^\s*Chapter\s+([IVXLCDM]+|\d+)\s*[–—-]\s*(.*?)_*$", re.I)
ITO_SUBSECTION = re.compile(r"^\s*\((\d+[A-Z]?)\)\s*(.*)$", re.I)
ITO_PARAGRAPH = re.compile(r"^\s*\(([a-z]{1,3})\)\s*(.*)$")


def _parse_income_tax_ordinance(data: bytes, filename: str, page_start: int = 1,
                               page_end: Optional[int] = None,
                               initial_chapter: Optional[str] = None,
                               initial_section: Optional[tuple] = None) -> Document:
    """Build a navigable chapter/section tree from the FBR's consolidated ITO PDF."""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("PDF support needs: pip install pypdf")

    reader = PdfReader(io.BytesIO(data))
    title = "Income Tax Ordinance, 2001"
    meta = {"Law": title, "Source Type": "Ordinance", "Amended Up To": "30 June 2026"}
    b = _Builder(title, meta)
    root = b.root
    root.path = title
    page_end = min(page_end or len(reader.pages), len(reader.pages))
    page_start = max(1, page_start)
    root.page_start, root.page_end = page_start, page_end
    chapter_nodes = {}
    current_chapter = root
    current_section = None
    current_content = root
    started = page_start > 1
    statute_started = page_start > 1
    pending_chapter = None

    def add_node(parent, level, typ, heading, text="", citation="", page=None):
        n = Node(f"~#{len(b.doc.nodes)}", "~", level, typ, heading, text,
                 parent.id, [], citation, (parent.path + " › " if parent.path else "") + heading,
                 page, page)
        parent.children.append(n.id)
        b.doc.nodes.append(n)
        return n

    if initial_chapter and page_start > 1:
        m = re.match(r"Chapter\s+([^ ]+)\s+—\s+(.*)", initial_chapter)
        if m:
            current_chapter = add_node(root, 1, "chapter", initial_chapter, page=page_start)
            chapter_nodes[m.group(1)] = current_chapter
    if initial_section and current_chapter is not root:
        citation, heading = initial_section
        current_section = add_node(current_chapter, 2, "section", heading, citation=citation, page=page_start)
        current_content = current_section

    def append_text(node, value, page):
        value = value.strip()
        if not value:
            return
        node.text = (node.text + " " + value).strip()
        node.page_end = page
        parent_id = node.parent
        while parent_id:
            parent = b.doc.nodes[int(parent_id.rsplit("#", 1)[1])]
            parent.page_end = page
            parent_id = parent.parent

    for page_no in range(page_start, page_end + 1):
        page = reader.pages[page_no - 1]
        raw = page.extract_text() or ""
        for line in raw.replace("\r", "").splitlines():
            s = " ".join(line.split())
            if not s or s in {".", "..", "…"}:
                continue
            if not statute_started:
                if re.search(r"following Ordinance promulgated by the President", s, re.I):
                    statute_started = True
                else:
                    continue

            chapter_heading = re.match(r"^CHAPTER\s+([IVXLCDM]+|\d+)\s*$", s, re.I)
            if chapter_heading:
                pending_chapter = chapter_heading.group(1).upper()
                continue
            if pending_chapter:
                name = s.strip(" _")
                if len(name) < 100 and not re.match(r"^\d+[A-Z]{0,3}\.\s", name, re.I):
                    key = pending_chapter
                    if key not in chapter_nodes:
                        current_chapter = add_node(root, 1, "chapter", f"Chapter {key} — {name}", page=page_no)
                        chapter_nodes[key] = current_chapter
                    else:
                        current_chapter = chapter_nodes[key]
                        current_chapter.page_end = page_no
                    pending_chapter = None
                    current_section = None
                    current_content = current_chapter
                    started = True
                    continue
                pending_chapter = None

            cm = ITO_CHAPTER.match(s)
            if cm:
                key, name = cm.group(1).upper(), cm.group(2).strip(" _")
                if key not in chapter_nodes:
                    current_chapter = add_node(root, 1, "chapter", f"Chapter {key} — {name}", page=page_no)
                    chapter_nodes[key] = current_chapter
                else:
                    current_chapter = chapter_nodes[key]
                    current_chapter.page_end = page_no
                current_section = None
                current_content = current_chapter
                started = True
                continue

            sm = ITO_SECTION.match(s)
            if sm and ("—" in s or "–" in s):
                secno, tail, rest = sm.group(1).upper(), sm.group(2).strip(), (sm.group(3) or "").strip()
                # Avoid treating numbered clauses or page artifacts as sections.
                if len(secno) <= 5 and tail and len(tail) < 180:
                    title_text = re.sub(r"\s+", " ", tail).strip(" .—–")
                    if title_text:
                        current_section = add_node(current_chapter, 2, "section",
                                                   f"Section {secno} — {title_text}", citation=secno, page=page_no)
                        current_content = current_section
                        if rest:
                            first_sub = ITO_SUBSECTION.match(rest)
                            first_para = ITO_PARAGRAPH.match(rest)
                            if first_sub:
                                label, content = first_sub.group(1), first_sub.group(2)
                                current_content = add_node(current_section, 3, "subsection", f"({label})", content,
                                                           f"{secno}({label})", page_no)
                            elif first_para:
                                label, content = first_para.group(1), first_para.group(2)
                                current_content = add_node(current_section, 3, "paragraph", f"({label})", content,
                                                           f"{secno}({label})", page_no)
                            else:
                                append_text(current_content, rest, page_no)
                        started = True
                        continue

            if not started:
                continue

            sub = ITO_SUBSECTION.match(s)
            if sub and current_section:
                label, rest = sub.group(1), sub.group(2)
                current_content = add_node(current_section, 3, "subsection", f"({label})", rest,
                                           f"{current_section.citation}({label})", page_no)
                continue
            para = ITO_PARAGRAPH.match(s)
            if para and current_section:
                label, rest = para.group(1), para.group(2)
                parent = current_content if current_content.type == "subsection" else current_section
                current_content = add_node(parent, parent.level + 1, "paragraph", f"({label})", rest,
                                           f"{parent.citation}({label})", page_no)
                continue
            if re.match(r"^(Provided that|Explanation|Provided further that)\b", s, re.I) and current_section:
                current_content = add_node(current_content, current_content.level + 1, "proviso", "Proviso", s,
                                           current_content.citation, page_no)
                continue

            # Ignore page numbers and recurring running headers; retain section body text.
            if re.fullmatch(r"\d+", s) or re.match(r"^(Chapter\s+)?[IVXLCDM]+\s*[–—-]\s*", s, re.I):
                continue
            append_text(current_content, s, page_no)

    b.doc.hint = "ito-2001"
    return b.doc


def _docx_text(data: bytes) -> str:
    try:
        import docx
    except ImportError:
        raise RuntimeError("DOCX support needs: pip install python-docx")
    out = []
    for p in docx.Document(io.BytesIO(data)).paragraphs:
        t, st = p.text.strip(), (p.style.name or "")
        if not t:
            continue
        if st == "Title":
            out.append("Title: " + t)
        elif st.startswith("Heading"):
            lvl = int(st.split()[-1]) if st.split()[-1].isdigit() else 1
            out.append("#" * lvl + " " + t)
        else:
            out.append(t)
    return "\n".join(out)


def parse_file(filename: str, data: bytes) -> List[Document]:
    base = filename.replace("\\", "/").split("/")[-1]
    name, ext = (base.rsplit(".", 1) + [""])[:2] if "." in base else (base, "txt")
    ext = ext.lower()
    if ext == "json":
        return from_json(json.loads(data.decode("utf-8-sig")), name)
    if ext == "pdf" and re.search(r"income.?tax.?ordinance|ordinanace", base, re.I):
        return [_parse_income_tax_ordinance(data, base)]
    text = _pdf_text(data) if ext == "pdf" else _docx_text(data) if ext == "docx" else data.decode("utf-8", "replace")
    return [parse_text(text, name)]
